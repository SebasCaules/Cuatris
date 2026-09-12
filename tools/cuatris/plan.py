"""Importador del plan de estudios S10-Rev23 y del vocabulario de sedes.

Fuentes y qué aporta cada una:

- `Plan S10-Rev23.xlsx`, hojas `Obligatorias` y `Electivas`: el recorte **vigente** del plan.
  De ahí salen `ciclo`, `cuatrimestre_sugerido`, `creditos_requeridos`, `correlativas` y
  `minors` de las 129 materias vigentes, y los mínimos de electivas (27) y de minor (14).
- `materias-carrera-info.html` (SGA): el listado histórico completo de 163 materias. De ahí
  salen `nombre` y `creditos` de todas, y los cuatro campos de las 34 que no están en el Excel.
- `oferta-carrera-info-titulos.html` (SGA): los créditos de cada título (147 / 192 / 243).

El módulo no depende de `canon.py`: serializa con `serializar_canonico`.

`cli.py` registra el subcomando llamando a `configurar_subcomando(subparsers)` y ejecuta
`args.funcion(args)` (alias `args.func`), que devuelve el código de salida del proceso.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from collections.abc import Iterable, Iterator, Sequence
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from bs4 import BeautifulSoup

CONTRATO = "1.0.0"
PLAN = "S10-Rev23"
CARRERA = "Ingeniería en Informática"

ENCABEZADO_MATERIAS = ("Materia", "Créditos", "Créditos requeridos", "Correlativas")

#: Encabezado de cada columna de minor en la hoja `Electivas`, en el orden del Excel.
#: El archivo del ITBA trae «Arqitectura de Software» con un error de tipeo; se corrige al
#: cargar y queda documentado en `docs/plan-de-estudios.md`.
MINORS_POR_ENCABEZADO: dict[str, tuple[str, str]] = {
    "Ciencia de Datos": ("CD", "Ciencia de Datos"),
    "Imágenes y Realidad Virtual": ("IRV", "Imágenes y Realidad Virtual"),
    "Inteligencia Artificial": ("IA", "Inteligencia Artificial"),
    "Arqitectura de Software": ("ARQ", "Arquitectura de Software"),
}

#: Orden en el que se emiten los minors, igual al de CONTRATO-v1.md §2.
ORDEN_MINORS = ("CD", "IA", "IRV", "ARQ")

#: Títulos que el contrato reconoce, indexados por el nombre del SGA sin acentos ni mayúsculas.
#: Los créditos **no** están aquí: se leen del HTML de títulos.
TITULOS_CONOCIDOS: dict[str, dict[str, object]] = {
    "analista en tecnologia informatica": {
        "id": "analista",
        "nombre": "Analista en Tecnología Informática",
        "tipo": "intermedio",
        "requiere_ciclos": ["basico"],
    },
    "bachiller en ingenieria": {
        "id": "bachiller",
        "nombre": "Bachiller en Ingeniería",
        "tipo": "intermedio",
        "requiere_ciclos": ["basico"],
    },
    "ingeniero/a en informatica": {
        "id": "ingeniero",
        "nombre": "Ingeniero/a en Informática",
        "tipo": "principal",
        "requiere_ciclos": ["basico", "profesional"],
        "requiere_electivas": True,
    },
}

#: Tipos de título del SGA que entran al plan. Las «Orientación» de 0 créditos son las viejas
#: orientaciones de la carrera, reemplazadas por los minors; no se cargan.
TIPOS_DE_TITULO = {"Intermedio": "intermedio", "Principal": "principal"}

#: Forma de cada `--creditos-decididos CODIGO=CREDITOS`. Cuando el Excel y el SGA declaran
#: créditos distintos para una materia, la importación aborta salvo que quien la corre resuelva
#: el conflicto materia por materia con este argumento: la elección la hace una persona en la
#: línea de comandos, no el importador.
RE_CREDITOS_DECIDIDOS = re.compile(r"^(\d{2}\.\d{2})=(\d+)$")

RE_CODIGO = re.compile(r"^(\d{2}\.\d{2})\s*-\s*(.+)$")
RE_CICLO = re.compile(r"^Ciclo\s+(Básico|Profesional)\b")
RE_ANIO_CUATRIMESTRE = re.compile(r"^Año\s+(\d+)\s*-\s*Cuatrimestre\s+(\d+)$")
RE_SUMA_CREDITOS = re.compile(r"suma\s+(\d+)\s+créditos")
RE_PIE_ELECTIVAS = re.compile(r"^Materias Electivas\b.*?(\d+)\s+créditos", re.DOTALL)
RE_MINIMO_MINOR = re.compile(r"mínimo de\s+(\d+)\s+créditos")
RE_MARCA_SEDE = re.compile(r"#-{2,}[^<\n]*")
RE_SEDE = re.compile(r"#-{2,}(?:&gt;|>)\s*([^<\n]+)")
RE_ID_SEDE = re.compile(r"^[a-z0-9_]+$")

CICLOS_EXCEL = {"Básico": "basico", "Profesional": "profesional"}


class ErrorPlan(Exception):
    """Algo en las fuentes no tiene la forma esperada. Nunca se produce un JSON a medias."""


# --------------------------------------------------------------------------- utilidades


def serializar_canonico(obj: object) -> str:
    """Forma canónica de CONTRATO-v1.md: claves ordenadas, indent 2, UTF-8 sin escapar."""
    return json.dumps(obj, ensure_ascii=False, indent=2, sort_keys=True) + "\n"


def escribir_json_canonico(ruta: Path, obj: object) -> None:
    ruta.parent.mkdir(parents=True, exist_ok=True)
    ruta.write_text(serializar_canonico(obj), encoding="utf-8", newline="\n")


def normalizar(texto: object) -> str:
    """Colapsa espacios (incluido U+00A0) y recorta. `None` pasa a cadena vacía."""
    if texto is None:
        return ""
    return " ".join(str(texto).replace("\xa0", " ").split())


def sin_acentos(texto: str) -> str:
    descompuesto = unicodedata.normalize("NFD", texto)
    return "".join(c for c in descompuesto if unicodedata.category(c) != "Mn")


def clave_titulo(nombre: str) -> str:
    return sin_acentos(normalizar(nombre)).lower()


def partir_codigo(celda: str) -> tuple[str, str]:
    """`'31.08 - Sistemas de Representación'` → `('31.08', 'Sistemas de Representación')`."""
    coincidencia = RE_CODIGO.match(celda)
    if not coincidencia:
        raise ErrorPlan(f"no parece una materia «código - nombre»: {celda!r}")
    return coincidencia.group(1), normalizar(coincidencia.group(2))


def partir_correlativas(celda: object) -> list[str]:
    """Las correlativas vienen separadas por espacios y U+00A0."""
    texto = normalizar(celda)
    if not texto:
        return []
    codigos = texto.split()
    for codigo in codigos:
        if not re.fullmatch(r"\d{2}\.\d{2}", codigo):
            raise ErrorPlan(f"correlativa con forma inesperada: {codigo!r} en {texto!r}")
    return codigos


def entero(valor: object, contexto: str) -> int:
    if isinstance(valor, bool) or not isinstance(valor, (int, float, str)):
        raise ErrorPlan(f"{contexto}: se esperaba un entero y vino {valor!r}")
    try:
        numero = int(str(valor).strip())
    except ValueError as error:
        raise ErrorPlan(f"{contexto}: se esperaba un entero y vino {valor!r}") from error
    if numero < 0:
        raise ErrorPlan(f"{contexto}: se esperaba un entero ≥ 0 y vino {numero}")
    return numero


# --------------------------------------------------------------------------- estructuras


@dataclass
class MateriaExcel:
    codigo: str
    nombre: str
    creditos: int
    creditos_requeridos: int | None
    correlativas: list[str]
    ciclo: str
    cuatrimestre_sugerido: int | None
    minors: list[str] = field(default_factory=list)


@dataclass
class MateriaSga:
    codigo: str
    nombre: str
    creditos: int
    creditos_requeridos: int
    correlativas: list[str]


@dataclass
class FuenteExcel:
    materias: dict[str, MateriaExcel]
    orden: list[str]
    creditos_por_ciclo: dict[str, int]
    creditos_electivas: int
    creditos_minimos_minor: int
    minors: list[dict[str, object]]


@dataclass
class FuenteSga:
    materias: dict[str, MateriaSga]
    orden: list[str]
    creditos_electivas: int


# --------------------------------------------------------------------------- Excel


def _filas_hoja(hoja) -> Iterator[Sequence[object]]:
    yield from hoja.iter_rows(values_only=True)


def _verificar_encabezado(fila: Sequence[object], esperado: Sequence[str], donde: str) -> None:
    leido = tuple(normalizar(c) for c in fila[: len(esperado)])
    if leido != tuple(esperado):
        raise ErrorPlan(
            f"{donde}: encabezado inesperado {list(leido)}; se esperaba {list(esperado)}"
        )


def leer_hoja_obligatorias(
    hoja,
) -> tuple[dict[str, MateriaExcel], list[str], dict[str, int], int, int]:
    """Lee la hoja `Obligatorias`: materias por ciclo/cuatrimestre y los mínimos del pie."""
    materias: dict[str, MateriaExcel] = {}
    orden: list[str] = []
    creditos_por_ciclo: dict[str, int] = {}
    creditos_electivas: int | None = None
    creditos_minimos_minor: int | None = None
    ciclo: str | None = None
    cuatrimestre: int | None = None
    en_pie = False

    for numero, fila in enumerate(_filas_hoja(hoja), start=1):
        primera = normalizar(fila[0] if fila else None)
        if not primera:
            continue

        pie = RE_PIE_ELECTIVAS.match(primera)
        if pie:
            en_pie = True
            creditos_electivas = int(pie.group(1))
            continue
        if en_pie:
            minimo = RE_MINIMO_MINOR.search(primera)
            if minimo:
                creditos_minimos_minor = int(minimo.group(1))
                continue
            # Debajo del pie solo hay prosa. Una materia o un bloque nuevo ahí abajo quedaría
            # fuera de todos los conteos de créditos, así que se aborta en vez de descartarlo.
            if (
                RE_CODIGO.match(primera)
                or RE_CICLO.match(primera)
                or RE_ANIO_CUATRIMESTRE.match(primera)
            ):
                raise ErrorPlan(
                    f"Obligatorias fila {numero}: {primera!r} aparece debajo del pie «Materias "
                    "Electivas»; el pie tiene que ser lo último de la hoja"
                )
            continue

        bloque = RE_CICLO.match(primera)
        if bloque:
            ciclo = CICLOS_EXCEL[bloque.group(1)]
            suma = RE_SUMA_CREDITOS.search(primera)
            if not suma:
                raise ErrorPlan(
                    f"Obligatorias fila {numero}: el bloque de ciclo no declara «suma N créditos»: "
                    f"{primera!r}"
                )
            creditos_por_ciclo[ciclo] = int(suma.group(1))
            cuatrimestre = None
            continue

        periodo = RE_ANIO_CUATRIMESTRE.match(primera)
        if periodo:
            if ciclo is None:
                raise ErrorPlan(
                    f"Obligatorias fila {numero}: bloque sin título de ciclo: {primera!r}"
                )
            anio, cuatri = int(periodo.group(1)), int(periodo.group(2))
            if cuatri not in (1, 2):
                raise ErrorPlan(f"Obligatorias fila {numero}: cuatrimestre inesperado {cuatri}")
            cuatrimestre = (anio - 1) * 2 + cuatri
            continue

        # Fila de encabezado: alcanza con que coincida una celda para exigir que coincidan todas.
        if set(normalizar(c) for c in fila[: len(ENCABEZADO_MATERIAS)]) & set(ENCABEZADO_MATERIAS):
            _verificar_encabezado(fila, ENCABEZADO_MATERIAS, f"Obligatorias fila {numero}")
            continue

        if not RE_CODIGO.match(primera):
            raise ErrorPlan(f"Obligatorias fila {numero}: fila que no se reconoce: {primera!r}")
        if ciclo is None or cuatrimestre is None:
            raise ErrorPlan(f"Obligatorias fila {numero}: materia fuera de un bloque: {primera!r}")

        codigo, nombre = partir_codigo(primera)
        if codigo in materias:
            raise ErrorPlan(f"Obligatorias fila {numero}: código repetido {codigo}")
        materias[codigo] = MateriaExcel(
            codigo=codigo,
            nombre=nombre,
            creditos=entero(fila[1], f"Obligatorias fila {numero} créditos"),
            creditos_requeridos=(
                None
                if fila[2] is None
                else entero(fila[2], f"Obligatorias fila {numero} requeridos")
            ),
            correlativas=partir_correlativas(fila[3] if len(fila) > 3 else None),
            ciclo=ciclo,
            cuatrimestre_sugerido=cuatrimestre,
        )
        orden.append(codigo)

    if creditos_electivas is None:
        raise ErrorPlan(
            "Obligatorias: falta el pie «Materias Electivas - Requisito: Completar N créditos»"
        )
    if creditos_minimos_minor is None:
        raise ErrorPlan("Obligatorias: falta la nota del mínimo de créditos de los minors")
    for esperado in ("basico", "profesional"):
        if esperado not in creditos_por_ciclo:
            raise ErrorPlan(f"Obligatorias: no aparece el bloque del ciclo {esperado}")
    return materias, orden, creditos_por_ciclo, creditos_electivas, creditos_minimos_minor


def leer_hoja_electivas(hoja) -> tuple[dict[str, MateriaExcel], list[str], list[dict[str, object]]]:
    """Lee la hoja `Electivas`: materias vigentes y la columna de minor a la que suman."""
    filas = list(_filas_hoja(hoja))
    if not filas:
        raise ErrorPlan("Electivas: la hoja está vacía")
    encabezado = [normalizar(c) for c in filas[0]]
    _verificar_encabezado(filas[0], ENCABEZADO_MATERIAS, "Electivas fila 1")

    columnas_minor: dict[int, str] = {}
    minors: list[dict[str, object]] = []
    primera_columna_de_minor = len(ENCABEZADO_MATERIAS)
    for indice, titulo in enumerate(
        encabezado[primera_columna_de_minor:], start=primera_columna_de_minor
    ):
        if not titulo:
            continue
        if titulo not in MINORS_POR_ENCABEZADO:
            raise ErrorPlan(f"Electivas: columna de minor desconocida {titulo!r}")
        sigla, nombre = MINORS_POR_ENCABEZADO[titulo]
        columnas_minor[indice] = sigla
        minors.append({"sigla": sigla, "nombre": nombre})
    faltantes = set(MINORS_POR_ENCABEZADO) - set(encabezado)
    if faltantes:
        raise ErrorPlan(f"Electivas: faltan columnas de minor: {sorted(faltantes)}")

    materias: dict[str, MateriaExcel] = {}
    orden: list[str] = []
    for numero, fila in enumerate(filas[1:], start=2):
        primera = normalizar(fila[0] if fila else None)
        if not primera:
            continue
        if not RE_CODIGO.match(primera):
            raise ErrorPlan(f"Electivas fila {numero}: fila que no se reconoce: {primera!r}")
        codigo, nombre = partir_codigo(primera)
        siglas = []
        for indice, sigla in columnas_minor.items():
            marca = normalizar(fila[indice]) if indice < len(fila) else ""
            if not marca:
                continue
            if marca.upper() != "X":
                raise ErrorPlan(
                    f"Electivas fila {numero}: marca de minor {sigla} desconocida en la columna "
                    f"{indice + 1}: {marca!r}; se esperaba «X», «x» o la celda vacía"
                )
            siglas.append(sigla)
        for indice in range(primera_columna_de_minor, len(fila)):
            if indice not in columnas_minor and normalizar(fila[indice]):
                raise ErrorPlan(
                    f"Electivas fila {numero}: la columna {indice + 1}, sin encabezado de minor, "
                    f"trae {normalizar(fila[indice])!r}"
                )
        materia = MateriaExcel(
            codigo=codigo,
            nombre=nombre,
            creditos=entero(fila[1], f"Electivas fila {numero} créditos"),
            creditos_requeridos=(
                None if fila[2] is None else entero(fila[2], f"Electivas fila {numero} requeridos")
            ),
            correlativas=partir_correlativas(fila[3] if len(fila) > 3 else None),
            ciclo="electiva",
            cuatrimestre_sugerido=None,
            minors=siglas,
        )
        anterior = materias.get(codigo)
        if anterior is None:
            materias[codigo] = materia
            orden.append(codigo)
            continue
        # El Excel repite 73.82 en dos filas idénticas salvo por la marca de minor.
        conflicto = (
            anterior.nombre != materia.nombre
            or anterior.creditos != materia.creditos
            or anterior.creditos_requeridos != materia.creditos_requeridos
            or anterior.correlativas != materia.correlativas
        )
        if conflicto:
            raise ErrorPlan(
                f"Electivas fila {numero}: {codigo} aparece dos veces con datos distintos "
                f"({anterior} vs {materia})"
            )
        unidos = set(anterior.minors) | set(materia.minors)
        anterior.minors = [s for s in ORDEN_MINORS if s in unidos]

    for materia in materias.values():
        materia.minors = [s for s in ORDEN_MINORS if s in materia.minors]
    return materias, orden, minors


def leer_excel(ruta: Path) -> FuenteExcel:
    libro = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    try:
        for hoja in ("Obligatorias", "Electivas"):
            if hoja not in libro.sheetnames:
                raise ErrorPlan(f"{ruta.name}: falta la hoja {hoja!r}; hay {libro.sheetnames}")
        (
            obligatorias,
            orden_obl,
            por_ciclo,
            creditos_electivas,
            minimo_minor,
        ) = leer_hoja_obligatorias(libro["Obligatorias"])
        electivas, orden_ele, minors = leer_hoja_electivas(libro["Electivas"])
    finally:
        libro.close()

    repetidas = set(obligatorias) & set(electivas)
    if repetidas:
        raise ErrorPlan(f"{ruta.name}: códigos en las dos hojas: {sorted(repetidas)}")

    materias = {**obligatorias, **electivas}
    suma_basico = sum(m.creditos for m in obligatorias.values() if m.ciclo == "basico")
    if suma_basico != por_ciclo["basico"]:
        raise ErrorPlan(
            f"{ruta.name}: el ciclo básico declara {por_ciclo['basico']} créditos y las materias "
            f"suman {suma_basico}"
        )
    suma_profesional = sum(m.creditos for m in obligatorias.values() if m.ciclo == "profesional")
    if suma_profesional + creditos_electivas != por_ciclo["profesional"]:
        raise ErrorPlan(
            f"{ruta.name}: el ciclo profesional declara {por_ciclo['profesional']} créditos y las "
            f"materias ({suma_profesional}) más las electivas ({creditos_electivas}) suman "
            f"{suma_profesional + creditos_electivas}"
        )
    return FuenteExcel(
        materias=materias,
        orden=orden_obl + orden_ele,
        creditos_por_ciclo=por_ciclo,
        creditos_electivas=creditos_electivas,
        creditos_minimos_minor=minimo_minor,
        minors=minors,
    )


# --------------------------------------------------------------------------- SGA


def _tablas_de_materias(sopa: BeautifulSoup) -> Iterator[list[list[str]]]:
    """Devuelve las tablas hoja cuyo encabezado es el de materias, ya como texto."""
    for tabla in sopa.find_all("table"):
        if tabla.find("table") is not None:
            continue  # es un envoltorio; las filas reales están en las tablas de adentro
        filas = [
            [normalizar(celda.get_text(" ", strip=True)) for celda in fila.find_all(["th", "td"])]
            for fila in tabla.find_all("tr")
        ]
        if not filas:
            continue
        tiene_materias = any(RE_CODIGO.match(f[0]) for f in filas[1:] if f)
        encabezado_ok = tuple(filas[0][: len(ENCABEZADO_MATERIAS)]) == ENCABEZADO_MATERIAS
        if not encabezado_ok:
            if tiene_materias:
                raise ErrorPlan(f"SGA: tabla con materias y encabezado inesperado {filas[0]}")
            continue
        yield filas


def leer_materias_sga(ruta: Path) -> FuenteSga:
    sopa = BeautifulSoup(ruta.read_text(encoding="utf-8"), "html.parser")
    materias: dict[str, MateriaSga] = {}
    orden: list[str] = []
    creditos_electivas: int | None = None

    for filas in _tablas_de_materias(sopa):
        for fila in filas[1:]:
            if not fila or not fila[0]:
                continue
            if len(fila) < len(ENCABEZADO_MATERIAS):
                raise ErrorPlan(f"SGA: fila con menos columnas de las esperadas: {fila}")
            if fila[0] == "Electivas":
                # Fila resumen del requisito de electivas, no es una materia.
                creditos_electivas = entero(fila[1], "SGA fila «Electivas» créditos")
                continue
            codigo, nombre = partir_codigo(fila[0])
            nueva = MateriaSga(
                codigo=codigo,
                nombre=nombre,
                creditos=entero(fila[1], f"SGA {codigo} créditos"),
                creditos_requeridos=entero(fila[2], f"SGA {codigo} créditos requeridos"),
                correlativas=partir_correlativas(fila[3]),
            )
            anterior = materias.get(codigo)
            if anterior is not None and anterior != nueva:
                raise ErrorPlan(f"SGA: {codigo} aparece dos veces con datos distintos")
            if anterior is None:
                materias[codigo] = nueva
                orden.append(codigo)

    if not materias:
        raise ErrorPlan(f"{ruta.name}: no se encontró ninguna tabla de materias")
    if creditos_electivas is None:
        raise ErrorPlan(f"{ruta.name}: no aparece la fila del requisito de electivas")
    return FuenteSga(materias=materias, orden=orden, creditos_electivas=creditos_electivas)


def leer_titulos_sga(ruta: Path) -> dict[str, int]:
    """Devuelve `id de título → créditos`, leídos de la tabla de títulos otorgados."""
    sopa = BeautifulSoup(ruta.read_text(encoding="utf-8"), "html.parser")
    esperado = ("Tipo", "Título", "Créditos requeridos")
    creditos: dict[str, int] = {}
    encontrada = False
    for tabla in sopa.find_all("table"):
        filas = [
            [normalizar(celda.get_text(" ", strip=True)) for celda in fila.find_all(["th", "td"])]
            for fila in tabla.find_all("tr")
        ]
        if not filas or tuple(filas[0][: len(esperado)]) != esperado:
            continue
        encontrada = True
        for fila in filas[1:]:
            if len(fila) < 3:
                raise ErrorPlan(f"títulos: fila con menos columnas de las esperadas: {fila}")
            tipo, nombre, valor = fila[0], fila[1], fila[2]
            if tipo not in TIPOS_DE_TITULO:
                continue  # «Orientación»: las viejas orientaciones, con 0 créditos
            datos = TITULOS_CONOCIDOS.get(clave_titulo(nombre))
            if datos is None:
                raise ErrorPlan(f"títulos: título desconocido {nombre!r} ({tipo})")
            if datos["tipo"] != TIPOS_DE_TITULO[tipo]:
                raise ErrorPlan(
                    f"títulos: {nombre!r} figura como {tipo!r} y el contrato lo tiene como "
                    f"{datos['tipo']!r}"
                )
            creditos[str(datos["id"])] = entero(valor, f"títulos {nombre}")
    if not encontrada:
        raise ErrorPlan(f"{ruta.name}: no se encontró la tabla de títulos otorgados")
    faltan = {str(d["id"]) for d in TITULOS_CONOCIDOS.values()} - set(creditos)
    if faltan:
        raise ErrorPlan(f"títulos: faltan en el HTML: {sorted(faltan)}")
    return creditos


# --------------------------------------------------------------------------- armado


def buscar_ciclos_correlativas(materias: Iterable[dict[str, object]]) -> list[list[str]]:
    """Devuelve los ciclos del grafo de correlativas; lista vacía si es acíclico."""
    hijos = {str(m["codigo"]): [str(c) for c in m["correlativas"]] for m in materias}
    estado: dict[str, int] = {}
    ciclos: list[list[str]] = []
    for inicio in hijos:
        if estado.get(inicio):
            continue
        pila: list[tuple[str, Iterator[str]]] = [(inicio, iter(hijos[inicio]))]
        camino = [inicio]
        estado[inicio] = 1
        while pila:
            nodo, pendientes = pila[-1]
            siguiente = next(pendientes, None)
            if siguiente is None:
                estado[nodo] = 2
                pila.pop()
                camino.pop()
                continue
            if estado.get(siguiente) == 1:
                ciclos.append(camino[camino.index(siguiente) :] + [siguiente])
            elif estado.get(siguiente, 0) == 0:
                estado[siguiente] = 1
                camino.append(siguiente)
                pila.append((siguiente, iter(hijos.get(siguiente, []))))
    return ciclos


def construir_plan(
    excel: FuenteExcel,
    sga: FuenteSga,
    creditos_titulos: dict[str, int],
    creditos_decididos: dict[str, int] | None = None,
) -> tuple[dict, dict]:
    """Cruza las tres fuentes y devuelve `(plan, reporte de diferencias)`.

    `creditos_decididos` resuelve, materia por materia, los casos en los que el Excel y el SGA
    declaran créditos distintos. Sin una entrada para cada una de esas materias la importación
    aborta: el importador nunca elige por su cuenta.
    """
    faltan_en_sga = [c for c in excel.orden if c not in sga.materias]
    if faltan_en_sga:
        raise ErrorPlan(
            "hay materias del Excel que no están en el listado del SGA: " + ", ".join(faltan_en_sga)
        )
    if excel.creditos_electivas != sga.creditos_electivas:
        raise ErrorPlan(
            f"el Excel pide {excel.creditos_electivas} créditos de electivas y el SGA "
            f"{sga.creditos_electivas}"
        )

    diferencias_creditos: list[dict[str, object]] = []
    diferencias_requeridos: list[dict[str, object]] = []
    diferencias_nombre: list[dict[str, object]] = []
    diferencias_correlativas: list[dict[str, object]] = []
    decididos = dict(creditos_decididos or {})
    decididos_usados: set[str] = set()
    sin_decidir: list[str] = []

    materias: list[dict[str, object]] = []
    for codigo in sga.orden:
        desde_sga = sga.materias[codigo]
        vigente = codigo in excel.materias
        desde_excel = excel.materias.get(codigo)

        creditos = desde_sga.creditos
        if desde_excel is not None and desde_excel.creditos != desde_sga.creditos:
            decidido = decididos.get(codigo)
            decididos_usados.add(codigo)
            if decidido is None:
                sin_decidir.append(
                    f"{codigo} ({desde_sga.nombre}): Excel {desde_excel.creditos} / "
                    f"SGA {desde_sga.creditos}"
                )
            elif decidido not in (desde_excel.creditos, desde_sga.creditos):
                raise ErrorPlan(
                    f"{codigo}: se decidieron {decidido} créditos, que no es ni el valor del "
                    f"Excel ({desde_excel.creditos}) ni el del SGA ({desde_sga.creditos})"
                )
            else:
                creditos = decidido
            diferencias_creditos.append(
                {
                    "codigo": codigo,
                    "excel": desde_excel.creditos,
                    "sga": desde_sga.creditos,
                    "decidido": decidido,
                }
            )
        if desde_excel is not None and desde_excel.nombre != desde_sga.nombre:
            diferencias_nombre.append(
                {"codigo": codigo, "excel": desde_excel.nombre, "sga": desde_sga.nombre}
            )

        if desde_excel is None:
            ciclo = "electiva"
            cuatrimestre = None
            requeridos = desde_sga.creditos_requeridos
            correlativas = list(desde_sga.correlativas)
            minors: list[str] = []
        else:
            ciclo = desde_excel.ciclo
            cuatrimestre = desde_excel.cuatrimestre_sugerido
            # El Excel manda para las vigentes; la celda vacía se completa con el SGA.
            requeridos = (
                desde_sga.creditos_requeridos
                if desde_excel.creditos_requeridos is None
                else desde_excel.creditos_requeridos
            )
            correlativas = list(desde_excel.correlativas)
            minors = list(desde_excel.minors)
            if requeridos != desde_sga.creditos_requeridos:
                diferencias_requeridos.append(
                    {"codigo": codigo, "excel": requeridos, "sga": desde_sga.creditos_requeridos}
                )
            if correlativas != desde_sga.correlativas:
                diferencias_correlativas.append(
                    {
                        "codigo": codigo,
                        "excel": correlativas,
                        "sga": list(desde_sga.correlativas),
                    }
                )

        materias.append(
            {
                "codigo": codigo,
                "nombre": desde_sga.nombre,
                "creditos": creditos,
                "ciclo": ciclo,
                "cuatrimestre_sugerido": cuatrimestre,
                "creditos_requeridos": requeridos,
                "correlativas": correlativas,
                "minors": minors,
                "vigente": vigente,
            }
        )

    if sin_decidir:
        raise ErrorPlan(
            "los créditos del Excel y del SGA difieren y nadie decidió cuál vale; agrega un "
            "--creditos-decididos CODIGO=CREDITOS por cada una: " + "; ".join(sin_decidir)
        )

    conocidos = {str(m["codigo"]) for m in materias}
    for materia in materias:
        for correlativa in materia["correlativas"]:
            if correlativa not in conocidos:
                raise ErrorPlan(
                    f"{materia['codigo']} declara la correlativa {correlativa}, que no está entre "
                    f"las {len(conocidos)} materias del plan"
                )
    ciclos = buscar_ciclos_correlativas(materias)
    if ciclos:
        detalle = "; ".join(" → ".join(c) for c in ciclos)
        raise ErrorPlan(f"el grafo de correlativas tiene ciclos: {detalle}")

    sobrantes = set(decididos) - decididos_usados
    if sobrantes:
        raise ErrorPlan(
            "--creditos-decididos nombra materias que no existen o cuyos créditos ya no difieren "
            "entre el Excel y el SGA: " + str(sorted(sobrantes))
        )

    titulos = []
    por_creditos = sorted(
        TITULOS_CONOCIDOS.values(), key=lambda d: creditos_titulos[str(d["id"])]
    )
    for datos in por_creditos:
        titulo: dict[str, object] = {
            "id": datos["id"],
            "nombre": datos["nombre"],
            "tipo": datos["tipo"],
            "creditos": creditos_titulos[str(datos["id"])],
            "requiere_ciclos": list(datos["requiere_ciclos"]),  # type: ignore[arg-type]
        }
        if datos.get("requiere_electivas"):
            titulo["requiere_electivas"] = excel.creditos_electivas
        titulos.append(titulo)

    if creditos_titulos["analista"] != excel.creditos_por_ciclo["basico"]:
        raise ErrorPlan(
            f"el título de Analista pide {creditos_titulos['analista']} créditos y el ciclo básico "
            f"del Excel suma {excel.creditos_por_ciclo['basico']}"
        )
    suma_ciclos = excel.creditos_por_ciclo["basico"] + excel.creditos_por_ciclo["profesional"]
    if creditos_titulos["ingeniero"] != suma_ciclos:
        raise ErrorPlan(
            f"el título de Ingeniero/a pide {creditos_titulos['ingeniero']} créditos y los dos "
            f"ciclos del Excel suman {suma_ciclos}"
        )

    minors = [
        {
            "sigla": sigla,
            "nombre": next(m["nombre"] for m in excel.minors if m["sigla"] == sigla),
            "creditos_minimos": excel.creditos_minimos_minor,
        }
        for sigla in ORDEN_MINORS
    ]
    siglas_usadas = {s for m in materias for s in m["minors"]}
    desconocidas = siglas_usadas - set(ORDEN_MINORS)
    if desconocidas:
        raise ErrorPlan(f"hay materias con minors desconocidos: {sorted(desconocidas)}")

    plan = {
        "contrato": CONTRATO,
        "plan": PLAN,
        "carrera": CARRERA,
        "titulos": titulos,
        "electivas": {"creditos_requeridos": excel.creditos_electivas},
        "minors": minors,
        "materias": materias,
    }
    reporte = {
        "total": len(materias),
        "vigentes": sum(1 for m in materias if m["vigente"]),
        "obligatorias_vigentes": sum(
            1 for m in materias if m["vigente"] and m["ciclo"] != "electiva"
        ),
        "electivas_vigentes": sum(
            1 for m in materias if m["vigente"] and m["ciclo"] == "electiva"
        ),
        "solo_sga": sum(1 for m in materias if not m["vigente"]),
        "diferencias_creditos": diferencias_creditos,
        "diferencias_creditos_requeridos": diferencias_requeridos,
        "diferencias_nombre": diferencias_nombre,
        "diferencias_correlativas": diferencias_correlativas,
    }
    return plan, reporte


def importar_plan(
    ruta_excel: Path,
    ruta_sga: Path,
    ruta_titulos: Path,
    creditos_decididos: dict[str, int] | None = None,
) -> tuple[dict, dict]:
    excel = leer_excel(ruta_excel)
    sga = leer_materias_sga(ruta_sga)
    creditos_titulos = leer_titulos_sga(ruta_titulos)
    return construir_plan(excel, sga, creditos_titulos, creditos_decididos)


# --------------------------------------------------------------------------- vocabulario


def extraer_sedes(html: str) -> list[str]:
    """Saca los nombres de sede de los tokens `Aula ITBA: <aula> #----> <sede>`.

    Toda marca que empiece con `#--` y no tenga esa forma aborta: si el SGA cambia el separador
    o deja de nombrar la sede, se entera quien corre el comando y no se publica un vocabulario
    incompleto en silencio.
    """
    nombres = []
    for marca in RE_MARCA_SEDE.finditer(html):
        fragmento = marca.group(0)
        coincidencia = RE_SEDE.match(fragmento)
        if not coincidencia:
            raise ErrorPlan(f"marca de sede con forma inesperada: {fragmento!r}")
        nombre = normalizar(coincidencia.group(1))
        if nombre.startswith("Sede "):
            nombre = nombre[len("Sede ") :]
        if not nombre:
            raise ErrorPlan(f"marca de sede sin nombre: {fragmento!r}")
        nombres.append(nombre)
    return nombres


def id_de_sede(nombre: str) -> str:
    identificador = re.sub(r"[^a-z0-9]+", "_", sin_acentos(nombre).lower()).strip("_")
    if not RE_ID_SEDE.fullmatch(identificador):
        raise ErrorPlan(f"no se pudo derivar un id válido para la sede {nombre!r}")
    return identificador


def construir_vocabulario(directorio_html: Path) -> dict:
    archivos = sorted(directorio_html.glob("*.html"))
    if not archivos:
        raise ErrorPlan(f"{directorio_html}: no hay archivos .html para leer")
    vistas: dict[str, str] = {}
    for archivo in archivos:
        for nombre in extraer_sedes(archivo.read_text(encoding="utf-8", errors="strict")):
            identificador = id_de_sede(nombre)
            anterior = vistas.get(identificador)
            if anterior is not None and anterior != nombre:
                raise ErrorPlan(
                    f"la sede {identificador!r} aparece como {anterior!r} y como {nombre!r}"
                )
            vistas[identificador] = nombre
    if not vistas:
        raise ErrorPlan(f"{directorio_html}: no se encontró ningún token de sede «#----> …»")
    sedes = [{"id": i, "nombre": n} for i, n in sorted(vistas.items())]
    return {"contrato": CONTRATO, "sedes": sedes}


# --------------------------------------------------------------------------- CLI


def par_creditos_decididos(texto: str) -> tuple[str, int]:
    """`'72.23=1'` → `('72.23', 1)`. Se usa como `type=` de `--creditos-decididos`."""
    coincidencia = RE_CREDITOS_DECIDIDOS.match(texto.strip())
    if not coincidencia:
        raise argparse.ArgumentTypeError(
            f"se esperaba «CODIGO=CREDITOS» (por ejemplo 72.23=1) y vino {texto!r}"
        )
    return coincidencia.group(1), int(coincidencia.group(2))


def _resumen(reporte: dict) -> str:
    lineas = [
        f"materias: {reporte['total']} ({reporte['vigentes']} vigentes = "
        f"{reporte['obligatorias_vigentes']} obligatorias + "
        f"{reporte['electivas_vigentes']} electivas; "
        f"{reporte['solo_sga']} solo en el SGA)",
    ]
    etiquetas = {
        "diferencias_creditos": "créditos (se deciden con --creditos-decididos)",
        "diferencias_creditos_requeridos": "créditos requeridos (manda el Excel)",
        "diferencias_nombre": "nombre (manda el SGA)",
        "diferencias_correlativas": "correlativas (manda el Excel)",
    }
    for clave, etiqueta in etiquetas.items():
        filas = reporte[clave]
        lineas.append(f"diferencias de {etiqueta}: {len(filas)}")
        for fila in filas:
            linea = f"  {fila['codigo']}: Excel {fila['excel']!r} / SGA {fila['sga']!r}"
            if "decidido" in fila:
                linea += f" → se usa {fila['decidido']!r}"
            lineas.append(linea)
    return "\n".join(lineas)


def _comando_importar(args: argparse.Namespace) -> int:
    decididos: dict[str, int] = {}
    for codigo, creditos in args.creditos_decididos or []:
        anterior = decididos.get(codigo)
        if anterior is not None and anterior != creditos:
            raise ErrorPlan(
                f"--creditos-decididos nombra {codigo} dos veces, con {anterior} y con {creditos}"
            )
        decididos[codigo] = creditos
    plan, reporte = importar_plan(
        Path(args.excel), Path(args.sga), Path(args.titulos), decididos
    )
    escribir_json_canonico(Path(args.salida), plan)
    print(_resumen(reporte), file=sys.stderr)
    print(f"escrito {args.salida}", file=sys.stderr)
    return 0


def _comando_vocabulario(args: argparse.Namespace) -> int:
    vocabulario = construir_vocabulario(Path(args.html))
    escribir_json_canonico(Path(args.salida), vocabulario)
    nombres = ", ".join(f"{s['id']} ({s['nombre']})" for s in vocabulario["sedes"])
    print(f"sedes encontradas: {nombres}", file=sys.stderr)
    print(f"escrito {args.salida}", file=sys.stderr)
    return 0


AYUDA = "Importa el plan de estudios y el vocabulario de sedes."

#: Acciones de `cuatris plan`, por si el parser se arma sin `set_defaults`.
ACCIONES = {"importar": _comando_importar, "vocabulario": _comando_vocabulario}


def configurar(parser: argparse.ArgumentParser) -> None:
    """Agrega las acciones de `cuatris plan` a un parser ya creado."""
    acciones = parser.add_subparsers(dest="accion", required=True)

    importar = acciones.add_parser("importar", help="Excel + listado del SGA → planes/<plan>.json")
    importar.add_argument("--excel", required=True, help="ruta al Plan S10-Rev23.xlsx")
    importar.add_argument("--sga", required=True, help="ruta a materias-carrera-info.html")
    importar.add_argument(
        "--titulos", required=True, help="ruta a oferta-carrera-info-titulos.html"
    )
    importar.add_argument(
        "--creditos-decididos",
        action="append",
        default=[],
        metavar="CODIGO=CREDITOS",
        type=par_creditos_decididos,
        help=(
            "resuelve a mano los créditos de una materia en la que el Excel y el SGA difieren "
            "(repetible); sin una entrada por cada diferencia la importación aborta"
        ),
    )
    importar.add_argument("--salida", required=True, help="archivo JSON a escribir")
    importar.set_defaults(funcion=_comando_importar, func=_comando_importar)

    vocabulario = acciones.add_parser(
        "vocabulario", help="HTML del SGA → vocabulario.json con las sedes"
    )
    vocabulario.add_argument("--html", required=True, help="directorio con los HTML del SGA")
    vocabulario.add_argument("--salida", required=True, help="archivo JSON a escribir")
    vocabulario.set_defaults(funcion=_comando_vocabulario, func=_comando_vocabulario)


def ejecutar(args: argparse.Namespace) -> int:
    """Corre la acción elegida y devuelve el código de salida."""
    return ACCIONES[args.accion](args)


def configurar_subcomando(subparsers) -> argparse.ArgumentParser:
    """Registra `cuatris plan` en los subparsers de `cli.py`."""
    plan = subparsers.add_parser("plan", help=AYUDA, description=AYUDA)
    plan.set_defaults(funcion=ejecutar, func=ejecutar)
    configurar(plan)
    return plan
